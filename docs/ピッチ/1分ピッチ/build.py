"""本選の1分ピッチ（ブースでデモを見せた直後に話す）を、pptx とプレビューPNGの両方に書き出す。

流れは「問いかけ → どんでん返し」。聞き手自身の薬箱から入り、
同じ商品の判定が在庫で変わる実証（百草丸）を2枚に割って見せる。
（体験談から入る案も作ったが、こちらに決めた）

同じ定義から両方を作るので、PNG を見れば pptx の配置を確かめられる。
挿絵は ../illust/（make.py の人物、props.py の薬箱と棚）。

  pip install python-pptx pillow
  python3 docs/ピッチ/1分ピッチ/build.py

プレビューPNGは、Noto Sans JP（Regular / Bold）が見つかったときだけ作る。
素材は一つ上の docs/ピッチ/assets を使う。
Googleスライドに取り込んだあとはそちらが正本。ここは叩き台を作り直すときに使う。
"""
import sys
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
A = HERE.parent / 'assets'
ARGS = [a for a in sys.argv[1:] if not a.startswith('--')]
OUT = Path(ARGS[0]) if ARGS else HERE
OUT.mkdir(parents=True, exist_ok=True)

W, H = 13.333, 7.5
FONT = 'Noto Sans JP'
DPI = 144

NAVY, INK, MUTED, FAINT = '1B2A4A', '0F172A', '475569', '94A3B8'
CANVAS, LINE, WHITE = 'F5F7FA', 'E2E8F0', 'FFFFFF'
RED, AMBER, BLUE, PALE = 'DC2626', 'C2410C', '2563EB', 'CBD5E1'

# プレビュー用のフォント。このフォルダか、OSのフォントフォルダから探す
FONT_DIR = next((d for d in (HERE, Path('C:/Windows/Fonts'), Path.home() / 'Library/Fonts', Path('/usr/share/fonts/opentype/noto'))
                 if (d / 'NotoSansJP-Bold.otf').exists()), None)


def font_file(bold):
    return FONT_DIR / ('NotoSansJP-Bold.otf' if bold else 'NotoSansJP-Regular.otf')


LEAD = 1.45  # Noto Sans JP の行送り（フォントサイズ比）


# ── 要素 ─────────────────────────────────────────────
def rect(x, y, w, h, fill, radius=0, line=None, label=None, size=18, color=INK, bold=True):
    """label を渡すと図形の中央に文字を置く（札・番号の丸）。"""
    return dict(kind='rect', x=x, y=y, w=w, h=h, fill=fill, radius=radius, line=line,
                label=label, size=size, color=color, bold=bold)


def text(x, y, w, h, body, size, color=INK, bold=False, align='left', spacing=1.0):
    return dict(kind='text', x=x, y=y, w=w, h=h, body=body, size=size, color=color,
                bold=bold, align=align, spacing=spacing)


def image(path, x, y, h=None, w=None):
    iw, ih = Image.open(A / path).size
    if h is None:
        h = w * ih / iw
    if w is None:
        w = h * iw / ih
    return dict(kind='image', path=path, x=x, y=y, w=w, h=h)


def phone(path, x, y, h):
    """スクリーンショットを白い枠に入れる。"""
    im = image(path, x, y, h=h)
    p = 0.09
    return [rect(x - p, y - p, im['w'] + 2 * p, h + 2 * p, WHITE, radius=0.22, line=LINE), im]


def brand(x, y, size=20, color=NAVY, logo='logo.png'):
    s = size / 72 * 1.5
    return [image(logo, x, y, h=s), text(x + s * 1.02, y + s * 0.5 - size * LEAD / 144, size / 72 * 4.5, size * LEAD / 72 + 0.05, 'Overlai', size, color)]



# ── 共通の部品 ──────────────────────────────────────
# 見る人はブースでデモを見た直後。1枚に映るのは7秒ほどなので、1枚に言葉は1つ。
# 判定の言葉はアプリの見出しと同じにする（lib/prompts.ts の headline の例）
AMBER_FILL = 'D97706'
X0 = 0.9            # 左右の余白
CW = W - 2 * X0     # 本文の幅

SLOGAN = '重ねる前に、重ねて見る。'
TEAM = '鈴鹿工業高等専門学校　井上 高志・濱田 圭太郎・杉本 隼都'


def cover():
    """表紙。名乗るだけ。スローガンは締めまで取っておく"""
    return dict(notes='鈴鹿高専の、Overlaiです。', items=[
        *brand(3.55, 2.55, size=96),
        text(0, 5.0, W, 0.5, 'ユメカタリ 学生生成AIコンテスト 2026｜開発部門', 16, MUTED, align='center'),
        text(0, 5.45, W, 0.5, TEAM, 16, MUTED, align='center'),
    ])


def statement(body, notes, size=54, color=NAVY, kicker=None, ill=None):
    """言葉だけのスライド。縦の中央に置く。
    ill を渡すと、言葉を左に寄せて右に挿絵を置く（絵は言葉の名詞そのものを描いたもの）"""
    lines = body.count('\n') + 1
    h = size * LEAD * lines / 72
    y = (H - h) / 2
    if ill:
        im = image(ill, 0, 0, w=4.3)
        im['x'], im['y'] = W - X0 - im['w'], (H - im['h']) / 2
        return dict(notes=notes, items=[text(X0, y, 7.2, h + 0.1, body, size, color, bold=True), im])
    items = [text(0, y, W, h + 0.1, body, size, color, bold=True, align='center')]
    if kicker:
        items.insert(0, text(0, y - 0.75, W, 0.5, kicker, 22, MUTED, bold=True, align='center'))
    return dict(notes=notes, items=items)


def head(body, size=40, y=0.95, color=NAVY):
    return text(X0, y, CW, size * LEAD / 72 + 0.1, body, size, color, bold=True)


# 家の薬箱の図。同じ図を2枚に割り、2枚目で百草丸を足す（段階ビルドアップ）
# 中身はシード（lib/seed.ts）から、名前を聞いて何か分かるものを選んだ
BOX_ITEMS = ['イブA錠', 'しっとり化粧水', 'ヒルドイドローション', 'ナイアシンアミド美容液']


def medbox(added=False):
    t = 3.0  # 図の上端。2枚で同じ位置にして、足した1行だけが動いて見えるようにする
    out = [
        # 店で撮った商品
        text(X0, t, 3.0, 0.4, '店で撮った商品', 18, MUTED, bold=True),
        rect(X0, t + 0.5, 3.0, 2.3, NAVY, radius=0.22),
        text(X0, t + 1.2, 3.0, 0.7, '御岳百草丸', 32, WHITE, bold=True, align='center'),
        text(X0, t + 1.9, 3.0, 0.4, '胃腸薬', 18, PALE, align='center'),
        text(3.95, t + 1.2, 0.6, 0.85, '×', 36, FAINT, bold=True, align='center'),
    ]
    # 家の薬箱
    bx, bw = 4.6, 4.75
    out += [
        text(bx, t, bw, 0.4, '家の薬箱', 18, MUTED, bold=True),
        rect(bx, t + 0.5, bw, 3.25 if added else 2.3, CANVAS, radius=0.22),
    ]
    for i, name in enumerate(BOX_ITEMS):
        out.append(text(bx + 0.35, t + 0.67 + i * 0.5, bw - 0.7, 0.45, name, 20, INK))
    if added:
        out += [
            rect(bx + 0.2, t + 2.78, bw - 0.4, 0.75, 'FFF7ED', radius=0.16, line=AMBER_FILL),
            text(bx + 0.35, t + 2.91, bw - 0.7, 0.5, '＋ 御岳百草丸', 22, AMBER, bold=True),
        ]
    # 判定（アプリの判定カードの見出しと同じ言葉・同じ色）
    color, word = (AMBER_FILL, '買わなくて\n大丈夫です') if added else (BLUE, '買っても\n問題なさそうです')
    out += [
        text(9.4, t + 1.2, 0.6, 0.85, '→', 36, FAINT, bold=True, align='center'),
        text(10.05, t, 2.4, 0.4, '判定', 18, MUTED, bold=True),
        rect(10.05, t + 0.5, 2.4, 2.3, color, radius=0.22),
        text(10.05, t + 1.05, 2.4, 1.3, word, 20, WHITE, bold=True, align='center'),
    ]
    return out


def guard_slide():
    """正直な注意点。医薬品の隣にあるアプリとして必ず問われる。
    実装：lib/verify.ts（裏取りと出典）、app/api/analyze/route.ts（相談の上書き）"""
    rows = [
        '成分をたどれない理由は消す',
        '出典を付けるのはAIではなくサーバー',
        '赤の判定には必ず相談先を添える',
    ]
    items = [head('AIの答えは、そのまま出しません。')]
    for i, r in enumerate(rows):
        y = 2.75 + i * 1.25
        items += [rect(X0, y + 0.1, 0.1, 0.62, NAVY), text(X0 + 0.45, y, CW - 0.5, 0.8, r, 32, INK, bold=True)]
    return dict(notes=(
        '医薬品の隣にあるので、AIの答えはそのまま出しません。'
        '根拠をたどれない理由は消し、赤のときは必ず相談先を案内します。'
    ), items=items)


def closing():
    """締め。まとめは置かず、スローガンで着地してブースへ戻す"""
    return dict(notes='買う前に、塗る前に、飲む前に。' + SLOGAN + 'Overlaiでした。', items=[
        text(0, 1.75, W, 1.2, SLOGAN, 60, NAVY, bold=True, align='center'),
        *brand(4.95, 3.55, size=48),
        text(X0, 6.55, 9, 0.4, TEAM, 15, MUTED),
        rect(11.28, 4.95, 1.4, 1.4, WHITE, radius=0.12, line=LINE),
        image('qr.png', 11.33, 5.0, h=1.3),
        text(10.88, 6.45, 2.2, 0.4, 'ブースで試せます', 13, MUTED, align='center'),
    ])


# ── スライドの並び ───────────────────────────────────
# 聞き手自身の薬箱から入り、「商品を調べれば分かる？」を、同じ商品の判定が変わる実証で覆す。
# 百草丸の実証は STATUS.md「在庫を変えると判定が変わることの実証」（実機で確認済み）
def episodes_pair():
    """2人分の実話を1枚に。絵は本文の名詞（迷っている本人）を描いている"""
    return dict(notes=(
        '井上は、風邪薬を買って帰ったら、似た薬がもう家にありました。'
        '濱田と杉本は、ステロイドを塗った肌に化粧水を使っていいのか、棚の前で迷いました。'
    ), items=[
        head('僕らの場合', size=32),
        image('people-2.png', 2.68, 1.95, h=2.5),
        text(X0, 4.7, 5.3, 0.4, '井上', 18, AMBER, bold=True, align='center'),
        text(X0, 5.15, 5.3, 1.0, '風邪薬を買って帰ったら\n似た薬がもう家にあった', 24, INK, bold=True, align='center'),
        image('people-1.png', 7.8, 1.95, h=2.5),
        text(6.85, 4.7, 5.6, 0.4, '濱田・杉本（ステロイドの塗り薬を使っている）', 18, BLUE, bold=True, align='center'),
        text(6.85, 5.15, 5.6, 1.0, '薬を塗った肌に\nこの化粧水、使っていい？', 24, INK, bold=True, align='center'),
    ])


SLIDES = [
    cover(),
    statement('家の薬箱の中身、\n全部言えますか？', '家の薬箱に何が入っているか、全部言えますか？', ill='medbox.png'),
    statement('ドラッグストアの\n棚の前でも？', 'では、ドラッグストアの棚の前で、思い出せますか？', ill='shelf.png'),
    episodes_pair(),
    statement('商品を調べれば分かる？', '商品を調べれば分かるでしょうか。', size=50),
    dict(notes='胃腸薬の百草丸を撮ると、青。', items=[head('御岳百草丸を撮ると…'), *medbox(False)]),
    dict(notes='家に百草丸を足して、同じ商品をもう一度撮ると、黄色。',
         items=[head('同じ商品なのに答えが変わった！'), *medbox(True)]),
    statement('答えは、商品ではなく\n家にある。', '答えは、商品ではなく家にあります。', size=58),
    guard_slide(),
    closing(),
]


# ── pptx ─────────────────────────────────────────────
def rgb(h):
    return RGBColor.from_string(h)


def set_font(run, size, color, bold):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, FONT
    f.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set('typeface', FONT)


def build_pptx(path, slides):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    for s in slides:
        sl = prs.slides.add_slide(blank)
        bg = sl.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(s.get('bg', WHITE))
        for it in s['items']:
            x, y, w, h = (Inches(it[k]) for k in ('x', 'y', 'w', 'h'))
            if it['kind'] == 'rect':
                shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if it['radius'] else MSO_SHAPE.RECTANGLE,
                                          x, y, w, h)
                if it['radius']:
                    shp.adjustments[0] = it['radius'] / min(it['w'], it['h'])
                shp.fill.solid()
                shp.fill.fore_color.rgb = rgb(it['fill'])
                if it['line']:
                    shp.line.color.rgb = rgb(it['line'])
                    shp.line.width = Pt(1)
                else:
                    shp.line.fill.background()
                shp.shadow.inherit = False
                tf = shp.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                if it['label']:
                    tf.word_wrap = False
                    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    r = p.add_run()
                    r.text = it['label']
                    set_font(r, it['size'], it['color'], it['bold'])
            elif it['kind'] == 'image':
                sl.shapes.add_picture(str(A / it['path']), x, y, w, h)
            else:
                tb = sl.shapes.add_textbox(x, y, w, h)
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                for i, line in enumerate(it['body'].split('\n')):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER}[it['align']]
                    p.line_spacing = it['spacing']
                    r = p.add_run()
                    r.text = line
                    set_font(r, it['size'], it['color'], it['bold'])
        sl.notes_slide.notes_text_frame.text = s['notes']
    prs.save(path)


# ── プレビュー ───────────────────────────────────────
def px(v):
    return round(v * DPI)


def render(i, s):
    img = Image.new('RGB', (px(W), px(H)), '#' + s.get('bg', WHITE))
    d = ImageDraw.Draw(img)
    for it in s['items']:
        box = [px(it['x']), px(it['y']), px(it['x'] + it['w']), px(it['y'] + it['h'])]
        if it['kind'] == 'rect':
            d.rounded_rectangle(box, radius=px(it['radius']), fill='#' + it['fill'],
                                outline=('#' + it['line']) if it['line'] else None, width=2)
            if it['label']:
                f = ImageFont.truetype(str(font_file(it['bold'])),
                                       round(it['size'] * DPI / 72))
                d.text(((box[0] + box[2]) / 2, (box[1] + box[3]) / 2), it['label'], font=f,
                       fill='#' + it['color'], anchor='mm')
        elif it['kind'] == 'image':
            im = Image.open(A / it['path']).convert('RGBA').resize((box[2] - box[0], box[3] - box[1]), Image.LANCZOS)
            img.paste(im, box[:2], im)
        else:
            f = ImageFont.truetype(str(font_file(it['bold'])),
                                   round(it['size'] * DPI / 72))
            pitch = it['size'] * LEAD * it['spacing'] * DPI / 72
            yy = box[1]
            for para in it['body'].split('\n'):
                # 箱の幅で折り返す（pptx 側の折り返しを真似る）
                lines, cur = [], ''
                for ch in para:
                    if f.getlength(cur + ch) > box[2] - box[0] and cur:
                        lines.append(cur)
                        cur = ch
                    else:
                        cur += ch
                lines.append(cur)
                for ln in lines:
                    lw = f.getlength(ln)
                    xx = box[0] if it['align'] == 'left' else (box[0] + box[2] - lw) / 2
                    d.text((xx, yy + pitch * 0.12), ln, font=f, fill='#' + it['color'])
                    yy += pitch
            if yy > box[3] + 2:
                print(f'  ! {i}: 文字が箱からはみ出しています → {it["body"][:16]}…')
    return img



def sheet(imgs, path):
    """全スライドを2列に並べた一覧"""
    s, g = 0.5, 20
    tw, th = int(px(W) * s), int(px(H) * s)
    rows = (len(imgs) + 1) // 2
    out = Image.new('RGB', (2 * tw + 3 * g, rows * th + (rows + 1) * g), '#888888')
    for k, im in enumerate(imgs):
        out.paste(im.resize((tw, th), Image.LANCZOS), (g + k % 2 * (tw + g), g + k // 2 * (th + g)))
    out.save(path)


build_pptx(OUT / 'Overlai_1分ピッチ.pptx', SLIDES)
if FONT_DIR:
    imgs = [render(f'スライド{i + 1}', s) for i, s in enumerate(SLIDES)]
    sheet(imgs, OUT / 'preview.png')
    if '--pages' in sys.argv:  # 1枚ずつの原寸も書き出す（目視の検品用）
        for i, im in enumerate(imgs):
            im.save(OUT / f'page-{i + 1:02d}.png')
total = sum(len(s['notes']) for s in SLIDES)
print(f'{len(SLIDES)}枚 / 台本 {total}字（約{total / 5:.0f}秒）')

