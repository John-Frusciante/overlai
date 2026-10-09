"""out/clips/ の動画を1枚に1本ずつ貼り、PowerPoint にまとめる。

動画はスライドを開くと自動で再生し、終わったら最後のコマで止まる。
ノートにはその場面のナレーションの原稿を入れる（話す内容の目安）。

  npm run slides:pptx … 動画の書き出しからまとめまで一度にやる
"""
import json
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu

HERE = Path(__file__).resolve().parent.parent
CLIPS = HERE / 'out' / 'clips'
OUT = HERE.parent / 'ピッチ' / '動画版' / 'Overlai_1分ピッチ_動画版.pptx'
LINES = {l['scene']: l['text'] for l in json.loads((HERE / 'narration.json').read_text())['lines']}

# PowerPoint で「開始：自動」にしたときと同じ形。スライドが始まると、すぐ再生する
AUTOPLAY = '''<p:timing {ns}><p:tnLst><p:par>
<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
 <p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
  <p:par><p:cTn id="3" fill="hold">
   <p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst>
   <p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
    <p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect">
     <p:stCondLst><p:cond delay="0"/></p:stCondLst>
     <p:childTnLst><p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr>
      <p:cTn id="6" dur="{ms}" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
     </p:cBhvr></p:cmd></p:childTnLst>
    </p:cTn></p:par>
   </p:childTnLst></p:cTn></p:par></p:childTnLst>
  </p:cTn></p:par>
 </p:childTnLst></p:cTn>
 <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
 <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
 </p:seq>
 <p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0">
  <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn>
  <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''

FADE = f'<p:transition {nsdecls("p")} spd="med"><p:fade/></p:transition>'


def seconds(mp4: Path) -> float:
    """mp4 の長さ（秒）。mvhd の timescale と duration から読む"""
    data = mp4.read_bytes()
    i = data.index(b'mvhd')
    version = data[i + 4]
    if version == 1:
        scale = int.from_bytes(data[i + 24:i + 28])
        dur = int.from_bytes(data[i + 28:i + 36])
    else:
        scale = int.from_bytes(data[i + 16:i + 20])
        dur = int.from_bytes(data[i + 20:i + 24])
    return dur / scale


prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]

clips = sorted(CLIPS.glob('*.mp4'))
for mp4 in clips:
    scene = mp4.stem.split('-', 1)[1]
    slide = prs.slides.add_slide(blank)
    movie = slide.shapes.add_movie(
        str(mp4), 0, 0, prs.slide_width, prs.slide_height,
        poster_frame_image=str(mp4.with_suffix('.png')), mime_type='video/mp4',
    )
    sld = slide._element
    sld.remove(sld.find(qn('p:timing')))
    sld.append(parse_xml(FADE))
    sld.append(parse_xml(AUTOPLAY.format(ns=nsdecls('p'), ms=round(seconds(mp4) * 1000), spid=movie.shape_id)))
    slide.notes_slide.notes_text_frame.text = LINES.get(scene, '')

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f'{len(clips)}枚 → {OUT}（{OUT.stat().st_size / 1e6:.1f}MB）')
